"""Document processing service for knowledge base documents.

Part of Frontier-Inspired Platform Simplification (Phase C).
Handles file upload, text extraction, chunking, and embedding generation.
"""
import os
import csv
import re
from typing import Optional
from uuid import UUID

from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.logging import get_logger
from app.config import get_settings
from app.models.knowledge_document import KnowledgeDocument
from app.models.knowledge_chunk import KnowledgeChunk

logger = get_logger(__name__)

# pgvector column is Vector(1536) — truncate oversized embeddings to fit
EXPECTED_EMBEDDING_DIM = 1536
settings = get_settings()

ALLOWED_FILE_TYPES = {
    "pdf", "docx", "txt", "csv", "md",
    # Code files (for project workspaces)
    "py", "js", "ts", "jsx", "tsx", "java", "go", "rs",
    "sql", "json", "yaml", "yml", "xml", "html", "css",
}
MAX_FILE_SIZE_BYTES = 50 * 1024 * 1024  # 50MB
CHUNK_SIZE = 1000  # tokens (approx 4 chars per token)
CHUNK_OVERLAP = 200  # tokens
CHUNK_SIZE_CHARS = CHUNK_SIZE * 4
CHUNK_OVERLAP_CHARS = CHUNK_OVERLAP * 4


def _sanitize_filename(filename: str) -> str:
    """Sanitize filename to prevent path traversal attacks."""
    # Strip any directory components
    filename = os.path.basename(filename)
    # Replace anything that isn't alphanumeric, dash, underscore, or dot
    return re.sub(r'[^\w\-_.]', '_', filename)


class DocumentProcessorService:
    """Service for uploading, parsing, chunking, and embedding documents."""

    def __init__(self, db: AsyncSession):
        self.db = db

    async def upload_document(
        self,
        kb_id: UUID,
        tenant_id: UUID,
        filename: str,
        file_content: bytes,
        file_size: int,
        user_id: Optional[UUID] = None,
    ) -> KnowledgeDocument:
        """Upload a document to a knowledge base.

        If a document with the same filename exists, creates a new version.
        """
        # Sanitize filename to prevent path traversal
        filename = _sanitize_filename(filename)

        # Validate file type
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if ext not in ALLOWED_FILE_TYPES:
            raise ValueError(f"Unsupported file type '.{ext}'. Allowed: {', '.join(ALLOWED_FILE_TYPES)}")

        # Validate file size
        if file_size > MAX_FILE_SIZE_BYTES:
            raise ValueError(f"File size {file_size} exceeds maximum {MAX_FILE_SIZE_BYTES // (1024*1024)}MB")

        # Check document count limit
        from app.services.knowledge_base import MAX_DOCS_PER_KB
        doc_count_result = await self.db.execute(
            select(func.count(KnowledgeDocument.id)).where(
                KnowledgeDocument.knowledge_base_id == kb_id,
                KnowledgeDocument.tenant_id == tenant_id,
            )
        )
        current_count = doc_count_result.scalar_one()
        if current_count >= MAX_DOCS_PER_KB:
            raise ValueError(f"Maximum of {MAX_DOCS_PER_KB} documents per knowledge base reached")

        # Determine version (check existing documents with same filename)
        version_result = await self.db.execute(
            select(func.max(KnowledgeDocument.version)).where(
                KnowledgeDocument.knowledge_base_id == kb_id,
                KnowledgeDocument.tenant_id == tenant_id,
                KnowledgeDocument.filename == filename,
            )
        )
        max_version = version_result.scalar_one()
        version = (max_version or 0) + 1

        # Store file locally (can be swapped to S3 later)
        storage_dir = os.path.join("data", "knowledge", str(tenant_id), str(kb_id))
        os.makedirs(storage_dir, exist_ok=True)
        storage_path = os.path.join(storage_dir, f"v{version}_{filename}")

        with open(storage_path, "wb") as f:
            f.write(file_content)

        # Create document record
        doc = KnowledgeDocument(
            tenant_id=tenant_id,
            knowledge_base_id=kb_id,
            filename=filename,
            file_type=ext,
            file_size_bytes=file_size,
            version=version,
            status="uploading",
            storage_path=storage_path,
            uploaded_by=user_id,
        )
        self.db.add(doc)
        await self.db.flush()

        logger.info(f"Uploaded document '{filename}' v{version} to KB {kb_id}")
        return doc

    async def process_document(self, document_id: UUID, tenant_id: UUID) -> None:
        """Process a document: parse, chunk, embed, and store chunks.

        This is designed to be called as a background job.
        """
        # Get document
        result = await self.db.execute(
            select(KnowledgeDocument).where(
                KnowledgeDocument.id == document_id,
                KnowledgeDocument.tenant_id == tenant_id,
            )
        )
        doc = result.scalar_one_or_none()
        if not doc:
            logger.error(f"Document {document_id} not found")
            return

        try:
            # Update status to processing
            doc.status = "processing"
            await self.db.flush()

            # Quick-fail if embedding API key is not configured
            # Check tenant LLM providers first, then fall back to env config
            from app.services.rag_service import RAGService
            rag_svc = RAGService(self.db)
            embedding_model, resolved_key, _ = await rag_svc._resolve_embedding_config()
            if not resolved_key:
                doc.status = "error"
                doc.error_message = f"No API key configured for embedding model '{embedding_model}'. Set one in Settings > LLM Providers or via environment variables."
                await self.db.flush()
                logger.error(f"Cannot process document {document_id}: no embedding API key")
                return

            # Parse document to text
            text = self._parse_document(doc.storage_path, doc.file_type)
            if not text.strip():
                doc.status = "error"
                doc.error_message = "No text content extracted from document"
                await self.db.flush()
                return

            # Chunk text using semantic or naive chunking
            from app.services.chunking import chunk_text as chunk_text_async
            chunk_results = await chunk_text_async(text)

            # Build chunk dicts with metadata
            chunks: list[dict] = []
            for cr in chunk_results:
                chunks.append({
                    "content": cr.content,
                    "metadata": {
                        "char_start": cr.start_index,
                        "char_end": cr.end_index,
                        "token_count": cr.token_count,
                        "filename": doc.filename,
                    },
                })

            # Contextual retrieval: prepend document context to each chunk
            if settings.USE_SEMANTIC_CHUNKING and chunk_results:
                try:
                    from app.services.contextual_retrieval import generate_chunk_contexts
                    contextualized = await generate_chunk_contexts(text, chunk_results)
                    for i, ctx_content in enumerate(contextualized):
                        chunks[i]["content"] = ctx_content
                except Exception as e:
                    logger.warning(f"Contextual retrieval failed, using raw chunks: {e}")

            # Generate embeddings for (possibly contextualized) chunk content
            embeddings = await self._generate_embeddings([c["content"] for c in chunks])

            # Validate embeddings — if all are None, mark as error
            non_null_count = sum(1 for e in embeddings if e is not None)
            if non_null_count == 0 and len(chunks) > 0:
                doc.status = "error"
                doc.error_message = "Embedding generation failed for all chunks. Check your embedding API key configuration."
                await self.db.flush()
                from app.services.knowledge_base import KnowledgeBaseService
                kb_service = KnowledgeBaseService(self.db)
                await kb_service.update_kb_counts(doc.knowledge_base_id, tenant_id)
                logger.error(f"All embeddings null for document {document_id}")
                return

            # Delete old chunks for this document (in case of re-processing)
            await self.db.execute(
                KnowledgeChunk.__table__.delete().where(
                    KnowledgeChunk.document_id == document_id
                )
            )

            # Normalize embedding dimensions to match pgvector column (1536)
            for idx, emb in enumerate(embeddings):
                if emb is not None and len(emb) != EXPECTED_EMBEDDING_DIM:
                    if len(emb) > EXPECTED_EMBEDDING_DIM:
                        # Truncate (Matryoshka-style models support this)
                        embeddings[idx] = emb[:EXPECTED_EMBEDDING_DIM]
                        logger.debug(f"Truncated embedding from {len(emb)} to {EXPECTED_EMBEDDING_DIM} dims")
                    else:
                        # Pad with zeros (rare, but prevents crashes)
                        embeddings[idx] = emb + [0.0] * (EXPECTED_EMBEDDING_DIM - len(emb))

            # Store chunks
            for i, chunk_data in enumerate(chunks):
                chunk = KnowledgeChunk(
                    tenant_id=tenant_id,
                    document_id=document_id,
                    knowledge_base_id=doc.knowledge_base_id,
                    content=chunk_data["content"],
                    chunk_index=i,
                    embedding=embeddings[i] if i < len(embeddings) else None,
                    chunk_metadata=chunk_data.get("metadata", {}),
                )
                self.db.add(chunk)

            # Update document status
            doc.status = "ready"
            doc.chunk_count = len(chunks)
            await self.db.flush()

            # Update KB aggregate counts
            from app.services.knowledge_base import KnowledgeBaseService
            kb_service = KnowledgeBaseService(self.db)
            await kb_service.update_kb_counts(doc.knowledge_base_id, tenant_id)

            logger.info(f"Processed document '{doc.filename}': {len(chunks)} chunks")

        except Exception as e:
            doc.status = "error"
            doc.error_message = str(e)[:500]
            await self.db.flush()
            logger.error(f"Error processing document {document_id}: {e}")
            # Don't re-raise — let the caller's db session commit the error status

    def _parse_document(self, storage_path: str, file_type: str) -> str:
        """Extract text from a document file."""
        if file_type == "pdf":
            return self._parse_pdf(storage_path)
        elif file_type == "docx":
            return self._parse_docx(storage_path)
        elif file_type == "txt" or file_type == "md":
            return self._parse_text(storage_path)
        elif file_type == "csv":
            return self._parse_csv(storage_path)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

    def _parse_pdf(self, path: str) -> str:
        """Extract text from PDF using PyMuPDF (fitz)."""
        import fitz  # pymupdf — already in requirements.txt

        doc = fitz.open(path)
        pages = []
        for i, page in enumerate(doc):
            page_text = page.get_text() or ""
            if page_text.strip():
                pages.append(f"[Page {i+1}]\n{page_text}")
        doc.close()
        return "\n\n".join(pages)

    def _parse_docx(self, path: str) -> str:
        """Extract text from DOCX using python-docx."""
        try:
            import docx
        except ImportError:
            raise ValueError("python-docx is required for DOCX support. Install with: pip install python-docx")

        doc = docx.Document(path)
        return "\n\n".join(para.text for para in doc.paragraphs if para.text.strip())

    def _parse_text(self, path: str) -> str:
        """Read plain text or markdown file."""
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()

    def _parse_csv(self, path: str) -> str:
        """Convert CSV to readable text."""
        rows = []
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            headers = next(reader, None)
            if headers:
                for row in reader:
                    row_text = ", ".join(f"{h}: {v}" for h, v in zip(headers, row) if v.strip())
                    if row_text:
                        rows.append(row_text)
        return "\n".join(rows)

    def _chunk_text(self, text: str) -> list[dict]:
        """Split text into overlapping chunks.

        Uses character-based chunking with ~4 chars per token approximation.
        """
        chunks = []
        start = 0
        text_len = len(text)

        while start < text_len:
            end = min(start + CHUNK_SIZE_CHARS, text_len)

            # Try to break at sentence/paragraph boundary
            if end < text_len:
                # Look for paragraph break
                newline_pos = text.rfind("\n\n", start + CHUNK_SIZE_CHARS // 2, end)
                if newline_pos > start:
                    end = newline_pos
                else:
                    # Look for sentence end
                    for sep in [". ", ".\n", "! ", "? "]:
                        sep_pos = text.rfind(sep, start + CHUNK_SIZE_CHARS // 2, end)
                        if sep_pos > start:
                            end = sep_pos + len(sep)
                            break

            chunk_text = text[start:end].strip()
            if chunk_text:
                chunks.append({
                    "content": chunk_text,
                    "metadata": {
                        "char_start": start,
                        "char_end": end,
                    },
                })

            # Advance with overlap
            start = end - CHUNK_OVERLAP_CHARS if end < text_len else text_len

        return chunks

    async def _generate_embeddings(self, texts: list[str]) -> list[list[float] | None]:
        """Generate embeddings using RAGService (shared path with Redis cache)."""
        if not texts:
            return []

        from app.services.rag_service import RAGService
        rag_service = RAGService(self.db)

        embeddings: list[list[float] | None] = []
        for text_input in texts:
            try:
                embedding = await rag_service.get_embedding(text_input)
                embeddings.append(embedding)
            except Exception as e:
                logger.warning(f"Embedding generation failed for chunk: {e}")
                embeddings.append(None)

        return embeddings
